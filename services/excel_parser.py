from pathlib import Path
from datetime import datetime
import openpyxl
import xlrd


_DATE_FORMATS = ["%d/%m/%y", "%d/%m/%Y", "%d-%m-%Y", "%d-%m-%y", "%Y-%m-%d"]


def _normalise_date(raw: str) -> str:
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return raw  # unknown format — store as-is


def _to_float(val) -> float:
    try:
        return float(val) if val is not None else 0.0
    except (ValueError, TypeError):
        return 0.0


def find_header_row(ws) -> int:
    for row_idx in range(1, ws.max_row + 1):
        if ws.cell(row_idx, 1).value == "Date":
            return row_idx
    raise ValueError("Header row with 'Date' column not found in worksheet")


def _find_header_row_xls(sheet) -> int:
    for row_idx in range(sheet.nrows):
        if sheet.cell_value(row_idx, 0) == "Date":
            return row_idx
    raise ValueError("Header row with 'Date' column not found in worksheet")


def parse_narration(narration) -> tuple[str, str]:
    if not narration:
        return "", ""
    narration = str(narration).strip()
    if narration.startswith("UPI-"):
        at_pos = narration.find("@")
        beneficiary = narration[4:at_pos] if at_pos != -1 else narration[4:]
        txn_note = narration.rsplit("-", 1)[-1].strip()
        return beneficiary, txn_note
    return narration, ""


def _parse_xlsx(file_path: Path) -> list[dict]:
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        header_row = find_header_row(ws)

        col_map = {}
        for col_idx in range(1, ws.max_column + 1):
            header = ws.cell(header_row, col_idx).value
            if header:
                col_map[str(header).strip()] = col_idx

        date_col = col_map.get("Date")
        narration_col = col_map.get("Narration")
        withdrawal_col = col_map.get("Withdrawal Amt.")
        deposit_col = col_map.get("Deposit Amt.")

        rows = []
        last_date_str = None
        for row_idx in range(header_row + 1, ws.max_row + 1):
            date_val = ws.cell(row_idx, date_col).value if date_col else None

            if date_val:
                if isinstance(date_val, datetime):
                    candidate = date_val.strftime("%Y-%m-%d")
                else:
                    raw = str(date_val).strip()
                    candidate = _normalise_date(raw)

                if any(ch.isdigit() for ch in candidate):
                    last_date_str = candidate

            narration_check = ws.cell(row_idx, narration_col).value if narration_col else None
            if not last_date_str or not narration_check:
                continue

            date_str = last_date_str
            beneficiary, txn_note = parse_narration(narration_check)

            withdraw_amount = _to_float(ws.cell(row_idx, withdrawal_col).value if withdrawal_col else None)
            deposit_amount = _to_float(ws.cell(row_idx, deposit_col).value if deposit_col else None)

            if withdraw_amount == 0.0 and deposit_amount == 0.0:
                continue

            rows.append({
                "date": date_str,
                "beneficiary": beneficiary,
                "txn_note": txn_note,
                "withdraw_amount": withdraw_amount,
                "deposit_amount": deposit_amount,
                "category": None,
            })

        return rows
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Could not read .xlsx file: {e}") from e


def _parse_xls(file_path: Path) -> list[dict]:
    try:
        wb = xlrd.open_workbook(str(file_path))
        ws = wb.sheet_by_index(0)
        header_row = _find_header_row_xls(ws)

        col_map = {}
        for col_idx in range(ws.ncols):
            header = ws.cell_value(header_row, col_idx)
            if header:
                col_map[str(header).strip()] = col_idx

        date_col = col_map.get("Date")
        narration_col = col_map.get("Narration")
        withdrawal_col = col_map.get("Withdrawal Amt.")
        deposit_col = col_map.get("Deposit Amt.")

        rows = []
        last_date_str = None
        for row_idx in range(header_row + 1, ws.nrows):
            date_val = ws.cell_value(row_idx, date_col) if date_col is not None else None
            cell_type = ws.cell_type(row_idx, date_col) if date_col is not None else None

            if date_val:
                if cell_type == xlrd.XL_CELL_DATE:
                    candidate = xlrd.xldate_as_datetime(date_val, wb.datemode).strftime("%Y-%m-%d")
                else:
                    raw = str(date_val).strip()
                    candidate = _normalise_date(raw)

                if any(ch.isdigit() for ch in candidate):
                    last_date_str = candidate

            narration_check = ws.cell_value(row_idx, narration_col) if narration_col is not None else None
            if not last_date_str or not narration_check:
                continue

            date_str = last_date_str
            beneficiary, txn_note = parse_narration(narration_check)

            withdraw_amount = _to_float(ws.cell_value(row_idx, withdrawal_col) if withdrawal_col is not None else None)
            deposit_amount = _to_float(ws.cell_value(row_idx, deposit_col) if deposit_col is not None else None)

            if withdraw_amount == 0.0 and deposit_amount == 0.0:
                continue

            rows.append({
                "date": date_str,
                "beneficiary": beneficiary,
                "txn_note": txn_note,
                "withdraw_amount": withdraw_amount,
                "deposit_amount": deposit_amount,
                "category": None,
            })

        return rows
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Could not read .xls file: {e}") from e


def parse_file(file_path: Path) -> list[dict]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".xlsx":
        return _parse_xlsx(file_path)
    elif suffix == ".xls":
        return _parse_xls(file_path)
    raise ValueError(f"Unsupported file format: {suffix}")
